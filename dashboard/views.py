import json
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth import get_user_model
from django.core.paginator import Paginator
from django.db.models import Avg, Count, F, Q, Sum
from django.db.models.functions import TruncDate, TruncMonth
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.template.loader import render_to_string
from django.utils import timezone
from django.views.decorators.http import require_POST

from categories.models import Category
from orders.models import Order, OrderItem
from products.models import Inventory, Product
from reviews.models import Review

User = get_user_model()


# ── Shared helpers ─────────────────────────────────────────────

def _json(data):
    """Serialize Decimal and date objects for Chart.js."""
    def default(o):
        if isinstance(o, Decimal):
            return float(o)
        if hasattr(o, 'isoformat'):
            return o.isoformat()
        raise TypeError(f'Not serializable: {type(o)}')
    return json.dumps(data, default=default)


def _fill_dates(qs_list, start, end):
    """Fill gaps in daily queryset results with zero-revenue entries."""
    lookup  = {row['date']: row for row in qs_list}
    result  = []
    current = start
    while current <= end:
        result.append(
            lookup.get(current, {'date': current, 'revenue': Decimal('0'), 'orders': 0})
        )
        current += timedelta(days=1)
    return result


PAID = Q(payment_status='paid')


# ══ Dashboard home ════════════════════════════════════════════

@staff_member_required
def home(request):
    now            = timezone.now()
    today          = now.date()
    month_start    = today.replace(day=1)
    last_m_start   = (month_start - timedelta(days=1)).replace(day=1)
    start_30       = today - timedelta(days=29)

    paid_qs = Order.objects.filter(PAID)

    # ── KPI values ─────────────────────────────────────────────
    total_revenue  = paid_qs.aggregate(t=Sum('total'))['t']               or Decimal('0')
    month_revenue  = paid_qs.filter(
                         created_at__date__gte=month_start
                     ).aggregate(t=Sum('total'))['t']                      or Decimal('0')
    last_m_revenue = paid_qs.filter(
                         created_at__date__gte=last_m_start,
                         created_at__date__lt=month_start
                     ).aggregate(t=Sum('total'))['t']                      or Decimal('0')

    total_orders   = Order.objects.count()
    month_orders   = Order.objects.filter(created_at__date__gte=month_start).count()

    total_customers = User.objects.filter(is_staff=False).count()
    new_customers   = User.objects.filter(
                          is_staff=False, date_joined__date__gte=month_start
                      ).count()

    total_products   = Product.objects.filter(is_active=True).count()
    low_stock_count  = Inventory.objects.filter(
                           quantity__gt=0, quantity__lte=F('low_stock_threshold')
                       ).count()
    out_of_stock_cnt = Inventory.objects.filter(quantity=0).count()

    revenue_change = (
        float((month_revenue - last_m_revenue) / last_m_revenue * 100)
        if last_m_revenue > 0
        else (100.0 if month_revenue > 0 else 0.0)
    )

    # ── Chart data ──────────────────────────────────────────────
    daily = list(
        paid_qs.filter(created_at__date__gte=start_30)
               .annotate(date=TruncDate('created_at'))
               .values('date')
               .annotate(revenue=Sum('total'), orders=Count('id'))
               .order_by('date')
    )
    daily_filled = _fill_dates(daily, start_30, today)

    status_counts = list(
        Order.objects.values('status')
                     .annotate(count=Count('id'))
                     .order_by('status')
    )

    # ── Recent orders + low stock ───────────────────────────────
    recent_orders = (
        Order.objects.select_related('user')
                     .order_by('-created_at')[:10]
    )
    low_stock = (
        Inventory.objects.filter(quantity__lte=F('low_stock_threshold'))
                         .select_related('product__category')
                         .order_by('quantity')[:10]
    )

    return render(request, 'dashboard/home.html', {
        'total_revenue':    total_revenue,
        'month_revenue':    month_revenue,
        'revenue_change':   round(revenue_change, 1),
        'total_orders':     total_orders,
        'month_orders':     month_orders,
        'total_customers':  total_customers,
        'new_customers':    new_customers,
        'total_products':   total_products,
        'low_stock_count':  low_stock_count,
        'out_of_stock_cnt': out_of_stock_cnt,
        'recent_orders':    recent_orders,
        'low_stock':        low_stock,
        # JSON for Chart.js
        'daily_dates':      _json([r['date'] for r in daily_filled]),
        'daily_revenue':    _json([r['revenue'] for r in daily_filled]),
        'daily_orders':     _json([r['orders'] for r in daily_filled]),
        'status_labels':    _json([s['status'].title() for s in status_counts]),
        'status_data':      _json([s['count'] for s in status_counts]),
    })


# ══ Sales analytics ═══════════════════════════════════════════

@staff_member_required
def sales_analytics(request):
    twelve_months_ago = (timezone.now() - timedelta(days=365)).date()

    monthly = list(
        Order.objects.filter(PAID, created_at__date__gte=twelve_months_ago)
                     .annotate(month=TruncMonth('created_at'))
                     .values('month')
                     .annotate(
                         revenue=Sum('total'),
                         orders=Count('id'),
                         avg_order=Avg('total'),
                     )
                     .order_by('month')
    )

    top_categories = list(
        OrderItem.objects.filter(order__payment_status='paid')
                         .values(name=F('product__category__name'))
                         .annotate(revenue=Sum(F('price') * F('quantity')))
                         .order_by('-revenue')[:8]
    )

    payment_methods = list(
        Order.objects.filter(PAID)
                     .values('payment__gateway')
                     .annotate(count=Count('id'), revenue=Sum('total'))
                     .order_by('-revenue')
    )

    summary = Order.objects.filter(PAID).aggregate(
        total_revenue  = Sum('total'),
        total_orders   = Count('id'),
        avg_order      = Avg('total'),
        total_discount = Sum('discount_amount'),
    )

    return render(request, 'dashboard/sales.html', {
        'summary':          summary,
        'payment_methods':  payment_methods,
        'monthly_labels':   _json([m['month'].strftime('%b %Y') for m in monthly]),
        'monthly_revenue':  _json([m['revenue']    for m in monthly]),
        'monthly_orders':   _json([m['orders']     for m in monthly]),
        'monthly_avg':      _json([float(m['avg_order'] or 0) for m in monthly]),
        'category_labels':  _json([c['name'] or 'Uncategorized' for c in top_categories]),
        'category_revenue': _json([c['revenue']    for c in top_categories]),
    })


# ══ Product analytics ═════════════════════════════════════════

@staff_member_required
def product_analytics(request):
    top_selling = list(
        OrderItem.objects.filter(order__payment_status='paid')
                         .values('product_name', 'product__slug')
                         .annotate(
                             total_qty     = Sum('quantity'),
                             total_revenue = Sum(F('price') * F('quantity')),
                         )
                         .order_by('-total_qty')[:10]
    )

    approved = Q(reviews__is_approved=True)
    most_reviewed = (
        Product.objects.filter(is_active=True)
                       .annotate(
                           review_count = Count('reviews', filter=approved),
                           avg_rating   = Avg('reviews__rating', filter=approved),
                       )
                       .filter(review_count__gt=0)
                       .order_by('-review_count')[:10]
    )

    category_distribution = (
        Category.objects.filter(level=0)
                        .annotate(
                            product_count=Count(
                                'products',
                                filter=Q(products__is_active=True)
                            )
                        )
                        .filter(product_count__gt=0)
                        .order_by('-product_count')[:8]
    )

    return render(request, 'dashboard/products.html', {
        'top_selling':      top_selling,
        'most_reviewed':    most_reviewed,
        'category_labels':  _json([c.name for c in category_distribution]),
        'category_counts':  _json([c.product_count for c in category_distribution]),
    })


# ══ Customer analytics ════════════════════════════════════════

@staff_member_required
def customer_analytics(request):
    twelve_months_ago = timezone.now() - timedelta(days=365)

    monthly_customers = list(
        User.objects.filter(is_staff=False, date_joined__gte=twelve_months_ago)
                    .annotate(month=TruncMonth('date_joined'))
                    .values('month')
                    .annotate(count=Count('id'))
                    .order_by('month')
    )

    top_customers = list(
        Order.objects.filter(PAID)
                     .values(
                         'user__id', 'user__email',
                         'user__first_name', 'user__last_name'
                     )
                     .annotate(
                         total_spend = Sum('total'),
                         order_count = Count('id'),
                         avg_order   = Avg('total'),
                     )
                     .order_by('-total_spend')[:10]
    )

    top_reviewers = list(
        Review.objects.values('user__email', 'user__first_name')
                      .annotate(review_count=Count('id'))
                      .order_by('-review_count')[:8]
    )

    return render(request, 'dashboard/customers.html', {
        'total_customers':   User.objects.filter(is_staff=False).count(),
        'active_customers':  User.objects.filter(is_staff=False, is_active=True).count(),
        'top_customers':     top_customers,
        'top_reviewers':     top_reviewers,
        'monthly_labels':    _json([m['month'].strftime('%b %Y') for m in monthly_customers]),
        'monthly_customers': _json([m['count'] for m in monthly_customers]),
    })


# ══ Inventory ════════════════════════════════════════════════

@staff_member_required
def inventory_view(request):
    base_qs = (
        Inventory.objects.select_related('product__category', 'product__brand')
                         .order_by('quantity')
    )

    q = request.GET.get('q', '').strip()
    if q:
        base_qs = base_qs.filter(product__name__icontains=q)

    out_of_stock = base_qs.filter(quantity=0)
    low_stock    = base_qs.filter(quantity__gt=0, quantity__lte=F('low_stock_threshold'))
    in_stock     = base_qs.filter(quantity__gt=F('low_stock_threshold'))

    return render(request, 'dashboard/inventory.html', {
        'out_of_stock': out_of_stock,
        'low_stock':    low_stock,
        'in_stock':     in_stock,
        'out_count':    out_of_stock.count(),
        'low_count':    low_stock.count(),
        'in_count':     in_stock.count(),
        'q':            q,
    })


@staff_member_required
@require_POST
def update_inventory(request, inventory_id):
    inv = get_object_or_404(Inventory, pk=inventory_id)
    try:
        qty = int(request.POST.get('quantity', -1))
        if qty < 0:
            raise ValueError
    except (ValueError, TypeError):
        return JsonResponse({'success': False, 'message': 'Invalid quantity.'}, status=400)

    inv.quantity = qty
    inv.save(update_fields=['quantity', 'updated_at'])

    if qty == 0:
        label, color = 'Out of Stock', '#f85149'
    elif qty <= inv.low_stock_threshold:
        label, color = f'Low Stock ({qty})', '#d29922'
    else:
        label, color = f'In Stock ({qty})', '#3fb950'

    return JsonResponse({'success': True, 'quantity': qty, 'label': label, 'color': color})


# ══ User management ═══════════════════════════════════════════

@staff_member_required
def user_management(request):
    users = User.objects.filter(is_staff=False).order_by('-date_joined')

    q = request.GET.get('q', '').strip()
    if q:
        users = users.filter(
            Q(email__icontains=q)      |
            Q(first_name__icontains=q) |
            Q(last_name__icontains=q)
        )

    users = users.annotate(
        order_count = Count('orders'),
        total_spend = Sum('orders__total', filter=Q(orders__payment_status='paid')),
    )

    paginator = Paginator(users, 25)
    page      = paginator.get_page(request.GET.get('page', 1))

    return render(request, 'dashboard/users.html', {
        'users': page,
        'q':     q,
        'total': paginator.count,
    })


@staff_member_required
@require_POST
def toggle_user(request, user_id):
    user           = get_object_or_404(User, pk=user_id, is_staff=False)
    user.is_active = not user.is_active
    user.save(update_fields=['is_active'])
    return JsonResponse({'success': True, 'is_active': user.is_active})


# ══ Reports page ══════════════════════════════════════════════

@staff_member_required
def reports_view(request):
    today = date.today()
    return render(request, 'dashboard/reports.html', {
        'today':       today.strftime('%Y-%m-%d'),
        'month_start': today.replace(day=1).strftime('%Y-%m-%d'),
    })


# ══ Exports ═══════════════════════════════════════════════════

def _parse_date_range(request):
    from datetime import datetime
    today = date.today()
    try:
        start = datetime.strptime(request.GET.get('from', ''), '%Y-%m-%d').date()
    except ValueError:
        start = today.replace(day=1)
    try:
        end = datetime.strptime(request.GET.get('to', ''), '%Y-%m-%d').date()
    except ValueError:
        end = today
    return start, end


@staff_member_required
def export_sales_pdf(request):
    start, end = _parse_date_range(request)

    orders = (
        Order.objects.filter(PAID, created_at__date__gte=start, created_at__date__lte=end)
                     .select_related('user')
                     .prefetch_related('items')
                     .order_by('-created_at')
    )
    summary = orders.aggregate(
        total_revenue  = Sum('total'),
        total_orders   = Count('id'),
        avg_order      = Avg('total'),
        total_discount = Sum('discount_amount'),
    )

    html = render_to_string('dashboard/reports/sales_pdf.html', {
        'orders': orders, 'summary': summary, 'start': start, 'end': end,
    }, request=request)

    try:
        from xhtml2pdf import pisa
        buffer      = BytesIO()
        pisa_status = pisa.CreatePDF(html, dest=buffer)
        if not pisa_status.err:
            buffer.seek(0)
            response = HttpResponse(buffer, content_type='application/pdf')
            response['Content-Disposition'] = (
                f'attachment; filename="ShopX-Sales-{start}-{end}.pdf"'
            )
            return response
    except ImportError:
        pass

    return HttpResponse(html)


@staff_member_required
def export_sales_excel(request):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    start, end = _parse_date_range(request)

    orders = (
        Order.objects.filter(PAID, created_at__date__gte=start, created_at__date__lte=end)
                     .select_related('user', 'payment')
                     .prefetch_related('items')
                     .order_by('-created_at')
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sales Report'

    hdr_font    = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    hdr_fill    = PatternFill(start_color='238636', end_color='238636', fill_type='solid')
    hdr_align   = Alignment(horizontal='center', vertical='center')
    center      = Alignment(horizontal='center')
    right       = Alignment(horizontal='right')
    money_fmt   = '#,##0.00'
    thin_border = Border(bottom=Side(style='thin', color='E8E8E8'))
    alt_fill    = PatternFill(start_color='F6FFF7', end_color='F6FFF7', fill_type='solid')
    total_font  = Font(name='Calibri', bold=True, color='238636', size=11)

    # Title
    ws.merge_cells('A1:H1')
    ws['A1'].value     = f'ShopX Sales Report  ·  {start:%d %b %Y} – {end:%d %b %Y}'
    ws['A1'].font      = Font(name='Calibri', bold=True, size=14, color='238636')
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 32

    # Summary
    summary = orders.aggregate(t=Sum('total'), n=Count('id'))
    ws['A2'] = f"Orders: {summary['n'] or 0}"
    ws['D2'] = f"Total Revenue: ₹{float(summary['t'] or 0):,.2f}"
    ws['A2'].font = ws['D2'].font = Font(name='Calibri', bold=True, color='555555')
    ws.row_dimensions[2].height = 20

    # Header row
    headers    = ['Order #', 'Date', 'Customer', 'Items', 'Subtotal', 'Discount', 'Total', 'Status']
    col_widths = [18, 14, 32, 8, 14, 12, 14, 14]

    for col, (hdr, width) in enumerate(zip(headers, col_widths), 1):
        cell           = ws.cell(row=4, column=col, value=hdr)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = hdr_align
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[4].height = 24

    # Data rows
    for row_num, order in enumerate(orders, 5):
        row_fill = alt_fill if row_num % 2 == 0 else None
        data = [
            order.order_number,
            order.created_at.strftime('%d %b %Y'),
            order.user.email if order.user else 'Guest',
            order.items.count(),
            float(order.subtotal),
            float(order.discount_amount),
            float(order.total),
            order.get_status_display(),
        ]
        for col, value in enumerate(data, 1):
            cell        = ws.cell(row=row_num, column=col, value=value)
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
            if col in (5, 6, 7):
                cell.number_format = money_fmt
                cell.alignment     = right
            elif col == 4:
                cell.alignment = center

    # Totals row
    last = orders.count() + 5
    ws.cell(row=last, column=1, value='TOTAL').font = total_font
    for col, field in [(5, 'subtotal'), (6, 'discount_amount'), (7, 'total')]:
        agg  = orders.aggregate(t=Sum(field))['t'] or Decimal('0')
        cell = ws.cell(row=last, column=col, value=float(agg))
        cell.font          = total_font
        cell.number_format = money_fmt
        cell.alignment     = right

    ws.freeze_panes = 'A5'

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = (
        f'attachment; filename="ShopX-Sales-{start}-{end}.xlsx"'
    )
    return response


@staff_member_required
def export_inventory_excel(request):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    inventory = (
        Inventory.objects.select_related('product__category', 'product__brand')
                         .order_by('quantity')
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Inventory'

    hdr_font  = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    hdr_fill  = PatternFill(start_color='238636', end_color='238636', fill_type='solid')
    center    = Alignment(horizontal='center')
    right     = Alignment(horizontal='right')

    ws.merge_cells('A1:G1')
    ws['A1'].value     = f'ShopX Inventory Report  ·  {date.today():%d %b %Y}'
    ws['A1'].font      = Font(name='Calibri', bold=True, size=14, color='238636')
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 30

    headers    = ['Product', 'Category', 'Brand', 'Price (₹)', 'Qty', 'Threshold', 'Status']
    col_widths = [42, 20, 18, 14, 10, 14, 16]

    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell           = ws.cell(row=3, column=col, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[3].height = 22

    STATUS_FG   = {'out': 'CC0000', 'low': '856404', 'good': '155724'}
    STATUS_BG   = {'out': 'FFCCCC', 'low': 'FFF3CC', 'good': 'CCFFCC'}

    for row_num, inv in enumerate(inventory, 4):
        if inv.quantity == 0:
            key, label = 'out',  'Out of Stock'
        elif inv.quantity <= inv.low_stock_threshold:
            key, label = 'low',  'Low Stock'
        else:
            key, label = 'good', 'In Stock'

        status_fill = PatternFill(
            start_color=STATUS_BG[key], end_color=STATUS_BG[key], fill_type='solid'
        )
        status_font = Font(name='Calibri', color=STATUS_FG[key], bold=True)

        data = [
            inv.product.name,
            inv.product.category.name if inv.product.category else '—',
            inv.product.brand.name    if inv.product.brand    else '—',
            float(inv.product.price),
            inv.quantity,
            inv.low_stock_threshold,
            label,
        ]
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col, value=value)
            if col == 4:
                cell.number_format = '#,##0.00'
                cell.alignment     = right
            elif col in (5, 6):
                cell.alignment = center
            elif col == 7:
                cell.fill      = status_fill
                cell.font      = status_font
                cell.alignment = center

    ws.freeze_panes = 'A4'

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = (
        f'attachment; filename="ShopX-Inventory-{date.today()}.xlsx"'
    )
    return response